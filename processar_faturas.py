#!/usr/bin/env python3
"""
processar_faturas.py — Processador de Faturas de Cartão de Crédito

Convenção de sinal:
  Positivo  → Débito   (gasto / compra)
  Negativo  → Crédito  (estorno / pagamento recebido)

Uso:
    python processar_faturas.py fatura1.pdf fatura2.xlsx
    python processar_faturas.py *.csv --saida resultado.xlsx --verbose
    python processar_faturas.py fatura.pdf --cartao "Itaú Visa Gold"
"""

import sys
import os
import re
import logging
import warnings
import argparse
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

# ──────────────────────────────────────────────────────────────────────────────
# Logging
# ──────────────────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(levelname)s: %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger(__name__)


# ──────────────────────────────────────────────────────────────────────────────
# Constantes
# ──────────────────────────────────────────────────────────────────────────────

COLUNAS_FINAIS = [
    "Data",
    "Descrição",
    "Valor (R$)",
    "Tipo",
    "Parcela",
    "Cartão",
    "Categoria",
    "Observações",
]

FORMATO_MOEDA = 'R$ #,##0.00'

CORES = {
    "cabecalho_bg":  "4472C4",
    "cabecalho_fg":  "FFFFFF",
    "linha_par":     "DCE6F1",
    "linha_impar":   "FFFFFF",
    "vermelho":      "C00000",
    "total_bg":      "FFF2CC",
    "titulo_resumo": "1F4E79",
}

# Palavras que indicam linhas de resumo (não transações)
SKIP_KEYWORDS = frozenset(
    [
        "total", "subtotal", "saldo", "pagamento mínimo", "pagamento minimo",
        "vencimento", "limite", "disponível", "disponivel", "fatura",
        "anterior", "encargo", "juros", "iof", "tarifa", "anuidade",
        "crédito rotativo", "taxa", "resumo", "extrato",
    ]
)


# ──────────────────────────────────────────────────────────────────────────────
# Utilitários de normalização
# ──────────────────────────────────────────────────────────────────────────────

_DATE_FORMATS = [
    "%d/%m/%Y", "%d/%m/%y",
    "%Y-%m-%d", "%d-%m-%Y", "%d-%m-%y",
    "%d.%m.%Y", "%d.%m.%y",
    "%Y/%m/%d",
    "%d %b %Y", "%d %b %y",
    "%d %B %Y", "%d %B %y",
    "%m/%d/%Y",
]

_DATA_RE = re.compile(r"^\d{1,2}[/\-.]\d{1,2}([/\-.]\d{2,4})?$")
_VALOR_RE = re.compile(r"^-?\s*R?\$?\s*\d[\d.,\s]*$")


def normalizar_data(valor: Any) -> Optional[str]:
    """Converte qualquer representação de data para DD/MM/AAAA."""
    if valor is None:
        return None
    if isinstance(valor, float) and pd.isna(valor):
        return None

    # Objeto datetime/date do Python ou pandas
    import datetime as _dt
    if isinstance(valor, (_dt.datetime, _dt.date)):
        return valor.strftime("%d/%m/%Y")
    try:
        ts = pd.Timestamp(valor)
        if not pd.isna(ts):
            return ts.strftime("%d/%m/%Y")
    except Exception:
        pass

    s = str(valor).strip()
    if not s or s.lower() in ("nan", "none", "nat", ""):
        return None

    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue

    # Formato DD/MM sem ano — assume o ano do arquivo ou atual
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})$", s)
    if m:
        year = getattr(normalizar_data, "_ano_referencia", datetime.now().year)
        try:
            return datetime(year, int(m.group(2)), int(m.group(1))).strftime("%d/%m/%Y")
        except ValueError:
            pass

    try:
        return pd.to_datetime(s, dayfirst=True).strftime("%d/%m/%Y")
    except Exception:
        logger.warning(f"  ⚠ Data não reconhecida: {s!r}")
        return s


def normalizar_valor(valor: Any) -> Optional[float]:
    """Converte qualquer representação de valor monetário BR/EN para float."""
    if valor is None:
        return None
    if isinstance(valor, float):
        return None if pd.isna(valor) else round(valor, 2)
    if isinstance(valor, int):
        return round(float(valor), 2)

    s = str(valor).strip()
    if not s or s.lower() in ("nan", "none", ""):
        return None

    # Detectar sinal
    negativo = bool(re.search(r"^[\-\(]|cr[eé]dito", s, re.IGNORECASE))
    s = re.sub(r"[R$\s()\-]", "", s)
    s = re.sub(r"(?i)crédito|credito|débito|debito|cr\b|db\b|c\b|d\b", "", s)

    has_dot = "." in s
    has_comma = "," in s

    if has_dot and has_comma:
        # Decide qual é separador decimal pela posição mais à direita
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")   # 1.234,56
        else:
            s = s.replace(",", "")                     # 1,234.56
    elif has_comma:
        s = s.replace(",", ".")                        # 1234,56
    # else: 1234.56 — mantém

    s = re.sub(r"[^\d.]", "", s)
    if not s:
        return None

    try:
        result = round(float(s), 2)
        return -result if negativo else result
    except ValueError:
        logger.warning(f"  ⚠ Valor não reconhecido: {valor!r}")
        return None


def detectar_parcela(descricao: Any) -> str:
    """Extrai 'N/M' de descrições como 'AMAZON 2/6' ou 'PARC 02/06'."""
    if not descricao or (isinstance(descricao, float) and pd.isna(descricao)):
        return ""
    s = str(descricao)
    patterns = [
        r"PARC(?:ELA)?\.?\s*(\d{1,2})[/\-](\d{1,2})",
        r"(\d{1,2})\s+DE\s+(\d{1,2})\b",
        r"\b(\d{1,2})/(\d{1,2})\b(?!\d)",
    ]
    for p in patterns:
        m = re.search(p, s, re.IGNORECASE)
        if m:
            return f"{int(m.group(1))}/{int(m.group(2))}"
    return ""


def determinar_tipo(valor: Any) -> str:
    """'Débito' para valores positivos, 'Crédito' para negativos."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    try:
        return "Crédito" if float(valor) < 0 else "Débito"
    except (TypeError, ValueError):
        return ""


# ──────────────────────────────────────────────────────────────────────────────
# Detecção de banco
# ──────────────────────────────────────────────────────────────────────────────

_BANCO_POR_NOME: List[Tuple[str, str]] = [
    ("nubank",       "Nubank"),
    ("itau",         "Itaú"),
    ("itaú",         "Itaú"),
    ("bradesco",     "Bradesco"),
    ("xp ",          "XP Investimentos"),
    ("xp_",          "XP Investimentos"),
    ("xpinvest",     "XP Investimentos"),
    ("c6bank",       "C6 Bank"),
    ("c6 bank",      "C6 Bank"),
    ("inter",        "Inter"),
    ("santander",    "Santander"),
    ("caixa",        "Caixa"),
    ("bancodobrasil","Banco do Brasil"),
    (" bb ",         "Banco do Brasil"),
    ("amex",         "American Express"),
    ("american",     "American Express"),
    ("next",         "Next"),
    ("pan",          "Banco Pan"),
    ("neon",         "Neon"),
    ("picpay",       "PicPay"),
]

_BANCO_POR_CONTEUDO: List[Tuple[str, str]] = [
    ("nubank",             "Nubank"),
    ("itaú unibanco",      "Itaú"),
    ("banco itaú",         "Itaú"),
    ("bradesco",           "Bradesco"),
    ("xp investimentos",   "XP Investimentos"),
    ("xp visa",            "XP Investimentos"),
    ("c6 bank",            "C6 Bank"),
    ("banco inter",        "Inter"),
    ("santander",          "Santander"),
    ("american express",   "American Express"),
    ("amex",               "American Express"),
    ("caixa economica",    "Caixa"),
    ("banco do brasil",    "Banco do Brasil"),
]


def detectar_banco_nome(caminho: Path) -> str:
    stem = caminho.stem.lower()
    for kw, banco in _BANCO_POR_NOME:
        if kw in stem:
            return banco
    return caminho.stem


def detectar_banco_conteudo(texto: str) -> Optional[str]:
    t = texto.lower()
    for kw, banco in _BANCO_POR_CONTEUDO:
        if kw in t:
            return banco
    return None


# ──────────────────────────────────────────────────────────────────────────────
# Mapeamento de colunas
# ──────────────────────────────────────────────────────────────────────────────

def mapear_colunas(colunas: List[Any]) -> Dict[str, str]:
    """
    Retorna {col_original: col_padrao} com base em palavras-chave.
    Primeira correspondência vence.
    """
    mapping: Dict[str, str] = {}
    cols = [(str(c), str(c).lower().strip()) for c in colunas]

    regras: List[Tuple[str, List[str]]] = [
        ("Data",       ["data", "date", "dt "]),
        ("Descrição",  ["descri", "lançamento", "lancamento", "históric",
                        "histor", "estabelec", "title", "description",
                        "merchant", "portador", "compra"]),
        ("Valor (R$)", ["valor", "amount", "value", "r$", "montante",
                        "total", "quantia"]),
        ("Parcela",    ["parcela", "parc", "installment"]),
        ("Categoria",  ["categ", "category", "tipo de gasto"]),
        ("Tipo",       ["tipo", "type", "natureza", "mov"]),
    ]

    for dest, keywords in regras:
        if dest in mapping.values():
            continue
        for orig, low in cols:
            if orig in mapping:
                continue
            if any(kw in low for kw in keywords):
                mapping[orig] = dest
                break

    return mapping


def inferir_linha_cabecalho(df: pd.DataFrame) -> Optional[int]:
    """Encontra índice da linha de cabeçalho em um DataFrame sem header."""
    kws = {"data", "date", "lançamento", "descrição", "descricao",
           "valor", "description", "amount", "estabelecimento", "histórico"}
    for i, row in df.iterrows():
        row_text = " ".join(str(v).lower() for v in row if v is not None)
        if sum(1 for kw in kws if kw in row_text) >= 2:
            return int(i)
    return None


# ──────────────────────────────────────────────────────────────────────────────
# Parsers
# ──────────────────────────────────────────────────────────────────────────────

class FaturaParser:
    """Classe base."""

    def __init__(self, caminho: str):
        self.caminho = Path(caminho)
        self.nome_cartao = detectar_banco_nome(self.caminho)
        self.senha: Optional[str] = None  # senha para PDFs protegidos
        # Tenta extrair ano de referência do nome do arquivo (ex: 20260311)
        m = re.search(r"(20\d{2})", self.caminho.stem)
        self._ano_ref: int = int(m.group(1)) if m else datetime.now().year

    def parsear(self) -> pd.DataFrame:
        raise NotImplementedError

    # ── normalização final ────────────────────────────────────────────────────

    def _finalizar(self, df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return pd.DataFrame(columns=COLUNAS_FINAIS)

        df = df.copy()

        if "Data" in df.columns:
            df["Data"] = df["Data"].apply(normalizar_data)

        if "Valor (R$)" in df.columns:
            df["Valor (R$)"] = df["Valor (R$)"].apply(
                lambda x: x if isinstance(x, float) and not pd.isna(x)
                else normalizar_valor(x)
            )

        if "Tipo" not in df.columns or df["Tipo"].isna().all() or (df["Tipo"] == "").all():
            df["Tipo"] = df["Valor (R$)"].apply(determinar_tipo)

        if "Parcela" not in df.columns:
            df["Parcela"] = df.get("Descrição", pd.Series(dtype=str)).apply(detectar_parcela)
        else:
            mask = df["Parcela"].isna() | (df["Parcela"] == "")
            df.loc[mask, "Parcela"] = (
                df.loc[mask, "Descrição"].apply(detectar_parcela)
                if "Descrição" in df.columns else ""
            )

        df["Cartão"] = self.nome_cartao

        for col in ("Categoria", "Observações"):
            if col not in df.columns:
                df[col] = ""

        for col in COLUNAS_FINAIS:
            if col not in df.columns:
                df[col] = ""

        df = df[COLUNAS_FINAIS].copy()

        # Remove linhas sem data ou sem valor
        df = df[df["Data"].notna() & (df["Data"] != "") & (~df["Data"].isin(["nan", "None"]))]
        df = df[df["Valor (R$)"].notna()]

        # Ordenar por data
        try:
            df.insert(0, "_sort", pd.to_datetime(df["Data"], format="%d/%m/%Y", errors="coerce"))
            df = df.sort_values("_sort").drop(columns=["_sort"])
        except Exception:
            pass

        return df.reset_index(drop=True)


# ── CSV ───────────────────────────────────────────────────────────────────────

class CSVParser(FaturaParser):
    """Parser para arquivos CSV (Nubank, Inter, C6, etc.)."""

    def parsear(self) -> pd.DataFrame:
        df = self._ler_csv()
        if df is None or df.empty:
            logger.warning(f"  ⚠ CSV vazio: {self.caminho.name}")
            return pd.DataFrame(columns=COLUNAS_FINAIS)

        # Detectar banco pelo conteúdo
        sample = " ".join(str(v) for v in df.values.flatten()[:200] if v)
        banco = detectar_banco_conteudo(sample)
        if banco:
            self.nome_cartao = banco

        # Tentar inferir cabeçalho se as colunas parecerem numéricas
        if not any(str(c).isalpha() for c in df.columns):
            idx = inferir_linha_cabecalho(df)
            if idx is not None:
                df.columns = df.iloc[idx]
                df = df.iloc[idx + 1:].reset_index(drop=True)

        mapeamento = mapear_colunas(list(df.columns))

        if "Data" not in mapeamento.values():
            logger.warning(f"  ⚠ Coluna de data não encontrada: {list(df.columns)}")

        result = pd.DataFrame()
        for orig, dest in mapeamento.items():
            if orig in df.columns:
                result[dest] = df[orig].values

        logger.info(f"  CSV: {len(result)} linhas | cartão={self.nome_cartao}")
        return self._finalizar(result)

    def _ler_csv(self) -> Optional[pd.DataFrame]:
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            for sep in (",", ";", "\t", "|"):
                try:
                    df = pd.read_csv(
                        self.caminho,
                        encoding=enc,
                        sep=sep,
                        dtype=str,
                        on_bad_lines="skip",
                    )
                    if len(df.columns) >= 2 and len(df) >= 1:
                        return df
                except Exception:
                    continue
        return None


# ── Excel ─────────────────────────────────────────────────────────────────────

class ExcelParser(FaturaParser):
    """Parser para XLSX/XLS (Itaú, XP, Bradesco, etc.)."""

    def parsear(self) -> pd.DataFrame:
        sheets = self._ler_sheets()
        if not sheets:
            logger.warning(f"  ⚠ Excel vazio: {self.caminho.name}")
            return pd.DataFrame(columns=COLUNAS_FINAIS)

        frames = []
        for sheet_name, df_raw in sheets.items():
            df = self._processar_sheet(df_raw, sheet_name)
            if df is not None and not df.empty:
                frames.append(df)

        if not frames:
            return pd.DataFrame(columns=COLUNAS_FINAIS)

        combined = pd.concat(frames, ignore_index=True)
        return self._finalizar(combined)

    def _ler_sheets(self) -> Optional[Dict[str, pd.DataFrame]]:
        try:
            xls = pd.ExcelFile(self.caminho)
            out = {}
            for name in xls.sheet_names:
                try:
                    out[name] = pd.read_excel(xls, sheet_name=name, header=None, dtype=str)
                except Exception as e:
                    logger.warning(f"  ⚠ Sheet '{name}': {e}")
            return out
        except Exception as e:
            logger.error(f"  ✗ Erro ao abrir Excel: {e}")
            return None

    def _processar_sheet(self, df: pd.DataFrame, sheet_name: str) -> Optional[pd.DataFrame]:
        if df.empty or df.shape[0] < 2:
            return None

        # Detectar banco
        sample = " ".join(str(v) for row in df.values[:20] for v in row if v)
        banco = detectar_banco_conteudo(sample)
        if banco:
            self.nome_cartao = banco

        # Encontrar linha de cabeçalho
        idx = inferir_linha_cabecalho(df)
        if idx is None:
            idx = self._encontrar_inicio_dados(df)
        if idx is None:
            return None

        headers = [str(h).strip() if h and str(h).strip() else f"col_{i}"
                   for i, h in enumerate(df.iloc[idx])]
        data = df.iloc[idx + 1:].copy()
        data.columns = headers
        data = data.dropna(how="all").reset_index(drop=True)

        if data.empty:
            return None

        mapeamento = mapear_colunas(headers)

        if not mapeamento:
            return None

        result = pd.DataFrame()
        for orig, dest in mapeamento.items():
            if orig in data.columns:
                result[dest] = data[orig].values

        # XP-specific: separar crédito/débito se houver coluna Tipo
        if "Tipo" in result.columns:
            def _ajustar_sinal_tipo(row: pd.Series) -> pd.Series:
                tipo = str(row.get("Tipo", "")).lower()
                val = row.get("Valor (R$)")
                if isinstance(val, str):
                    val = normalizar_valor(val)
                if val is not None and "crédito" in tipo or "credito" in tipo:
                    row["Valor (R$)"] = -abs(val) if val else val
                return row

            result = result.apply(_ajustar_sinal_tipo, axis=1)

        logger.info(f"  Sheet '{sheet_name}': {len(result)} linhas | cartão={self.nome_cartao}")
        return result

    def _encontrar_inicio_dados(self, df: pd.DataFrame) -> Optional[int]:
        """Primeira linha que contém uma data → volta uma para pegar o cabeçalho."""
        for i, row in df.iterrows():
            vals = [str(v).strip() for v in row if v and str(v).strip()]
            if any(_DATA_RE.match(v) for v in vals):
                return max(0, int(i) - 1)
        return None


# ── PDF ───────────────────────────────────────────────────────────────────────

class PDFParser(FaturaParser):
    """
    Parser para PDFs digitais usando pdfplumber.

    Estratégia:
    1. Extração de tabelas (melhor qualidade)
    2. Parse de texto linha-por-linha
    3. OCR via pytesseract (se instalado)
    """

    def parsear(self) -> pd.DataFrame:
        # Expõe o ano de referência para normalizar_data()
        normalizar_data._ano_referencia = self._ano_ref  # type: ignore[attr-defined]
        try:
            open_kwargs: Dict[str, Any] = {}
            if self.senha:
                open_kwargs["password"] = self.senha
            with pdfplumber.open(self.caminho, **open_kwargs) as pdf:
                texto_total = "\n".join(p.extract_text() or "" for p in pdf.pages)
                banco = detectar_banco_conteudo(texto_total)
                if banco:
                    self.nome_cartao = banco
                logger.info(f"  PDF: {len(pdf.pages)} pág. | banco={self.nome_cartao}")

                # ── Estratégia 1: tabelas ──────────────────────────────────
                # Normaliza cada tabela individualmente antes de concatenar
                # para evitar erros de colunas incompatíveis entre tabelas.
                frames_tabelas = []
                for page in pdf.pages:
                    for table in page.extract_tables():
                        df_raw = self._tabela_para_df(table)
                        if df_raw is None or df_raw.empty:
                            continue
                        mapa = mapear_colunas(list(df_raw.columns))
                        # Só aceita tabelas que tenham Data E Valor
                        if "Data" not in mapa.values() or "Valor (R$)" not in mapa.values():
                            continue
                        norma = self._normalizar_tabela(df_raw)
                        if not norma.empty:
                            frames_tabelas.append(norma)

                if frames_tabelas:
                    raw = pd.concat(frames_tabelas, ignore_index=True)
                    if not raw.empty:
                        return self._finalizar(raw)

                # ── Estratégia 2: texto ────────────────────────────────────
                logger.info("  → nenhuma tabela, tentando parse de texto…")
                result = self._parsear_texto(texto_total)
                if not result.empty:
                    return self._finalizar(result)

                # ── Estratégia 3: OCR ──────────────────────────────────────
                try:
                    result = self._ocr()
                    if not result.empty:
                        return self._finalizar(result)
                except ImportError:
                    logger.warning("  ⚠ pytesseract/pdf2image não instalados — OCR indisponível")

                logger.warning(f"  ⚠ Nenhuma transação encontrada em {self.caminho.name}")
                return pd.DataFrame(columns=COLUNAS_FINAIS)

        except Exception as e:
            msg = str(e)
            if "PDFPasswordIncorrect" in type(e).__name__ or "password" in msg.lower():
                logger.error(
                    f"  ✗ PDF protegido por senha: {self.caminho.name}\n"
                    "     Use --senha SUASENHA (geralmente CPF ou data de nascimento)"
                )
            else:
                logger.error(f"  ✗ Erro ao processar PDF: {e}")
            return pd.DataFrame(columns=COLUNAS_FINAIS)

    # ── helpers de tabela ─────────────────────────────────────────────────────

    def _tabela_para_df(self, table: List) -> Optional[pd.DataFrame]:
        if not table or len(table) < 2:
            return None
        clean = [[str(c).strip() if c else "" for c in row] for row in table if any(row)]
        if len(clean) < 2:
            return None
        # Rejeita tabela se a primeira célula do cabeçalho for um blob de texto longo
        if len(clean[0]) == 1 and len(clean[0][0]) > 100:
            return None
        try:
            # Garante nomes de coluna únicos
            headers = clean[0]
            seen: Dict[str, int] = {}
            unique_headers = []
            for h in headers:
                if h in seen:
                    seen[h] += 1
                    unique_headers.append(f"{h}_{seen[h]}")
                else:
                    seen[h] = 0
                    unique_headers.append(h)
            df = pd.DataFrame(clean[1:], columns=unique_headers)
        except Exception:
            df = pd.DataFrame(clean)
        # Filtra linhas de subtotal/resumo
        if len(df.columns) > 0:
            first_col = df.iloc[:, 0].str.lower().str.strip()
            skip = first_col.isin(SKIP_KEYWORDS) | first_col.str.startswith("subtotal")
            df = df[~skip]
        return df if not df.empty else None

    def _normalizar_tabela(self, df: pd.DataFrame) -> pd.DataFrame:
        mapeamento = mapear_colunas(list(df.columns))
        if not mapeamento:
            return pd.DataFrame()
        result = pd.DataFrame()
        for orig, dest in mapeamento.items():
            if orig in df.columns:
                result[dest] = df[orig].values
        return result

    # ── helpers de texto ──────────────────────────────────────────────────────

    def _parsear_texto(self, texto: str) -> pd.DataFrame:
        rows = (
            self._regex_data_desc_valor(texto)
            or self._itau_texto(texto)
            or self._xp_texto(texto)
            or self._generico_texto(texto)
        )
        return pd.DataFrame(rows) if rows else pd.DataFrame()

    def _regex_data_desc_valor(self, texto: str) -> List[Dict]:
        """Padrão canônico: DATA DESCRIÇÃO VALOR no final da linha."""
        pattern = re.compile(
            r"^(\d{1,2}[/\-.]\d{1,2}(?:[/\-.]\d{2,4})?)"
            r"\s+(.+?)\s+"
            r"(-?\s*R?\$?\s*[\d.,]+)\s*$",
            re.MULTILINE,
        )
        rows = []
        for m in pattern.finditer(texto):
            data = m.group(1).strip()
            desc = re.sub(r"\s+", " ", m.group(2)).strip()
            val_str = m.group(3).strip()
            if any(kw in desc.lower() for kw in SKIP_KEYWORDS):
                continue
            val = normalizar_valor(val_str)
            if val is not None and len(desc) >= 3:
                rows.append({"Data": data, "Descrição": desc, "Valor (R$)": val})
        return rows

    def _itau_texto(self, texto: str) -> List[Dict]:
        """Parse específico para texto de fatura Itaú (DD/MM + desc + valor)."""
        pattern = re.compile(
            r"(\d{2}/\d{2})\s+"
            r"([A-Za-záéíóúÁÉÍÓÚàãõâêôçÀÃÕÂÊÔÇ\w\s\.\-\*/&'#]+?)\s+"
            r"([\d.,]+)\s*$",
            re.MULTILINE,
        )
        year = str(datetime.now().year)
        rows = []
        for m in pattern.finditer(texto):
            desc = re.sub(r"\s+", " ", m.group(2)).strip()
            if any(kw in desc.lower() for kw in SKIP_KEYWORDS) or len(desc) < 3:
                continue
            val = normalizar_valor(m.group(3))
            if val is not None:
                rows.append({
                    "Data": f"{m.group(1)}/{year}",
                    "Descrição": desc,
                    "Valor (R$)": val,
                })
        return rows

    def _xp_texto(self, texto: str) -> List[Dict]:
        """Parse específico para fatura XP Investimentos."""
        # XP usa formato: DD/MM/AAAA | ESTABELECIMENTO | PARCELA | VALOR
        pattern = re.compile(
            r"(\d{2}/\d{2}/\d{4})\s+"
            r"(.+?)\s+"
            r"(?:(\d{1,2}/\d{1,2})\s+)?"
            r"([\d.,]+)\s*$",
            re.MULTILINE,
        )
        rows = []
        for m in pattern.finditer(texto):
            desc = re.sub(r"\s+", " ", m.group(2)).strip()
            if any(kw in desc.lower() for kw in SKIP_KEYWORDS) or len(desc) < 3:
                continue
            val = normalizar_valor(m.group(4))
            if val is not None:
                rows.append({
                    "Data": m.group(1),
                    "Descrição": desc,
                    "Parcela": m.group(3) or "",
                    "Valor (R$)": val,
                })
        return rows

    def _generico_texto(self, texto: str) -> List[Dict]:
        """Fallback: procura qualquer linha com data + valor."""
        rows = []
        for linha in texto.splitlines():
            linha = linha.strip()
            if len(linha) < 8:
                continue
            dm = re.search(r"\b(\d{1,2}[/\-.]\d{1,2}(?:[/\-.]\d{2,4})?)\b", linha)
            if not dm:
                continue
            vals = re.findall(r"-?\s*R?\$?\s*\d[\d.,]{1,}", linha)
            if not vals:
                continue
            val = normalizar_valor(vals[-1])
            if val is None:
                continue
            start = dm.end()
            end = linha.rfind(vals[-1])
            desc = (linha[start:end] if end > start else linha[dm.end():]).strip()
            desc = re.sub(r"\s+", " ", desc)
            if any(kw in desc.lower() for kw in SKIP_KEYWORDS) or len(desc) < 2:
                continue
            rows.append({"Data": dm.group(1), "Descrição": desc, "Valor (R$)": val})
        return rows

    def _ocr(self) -> pd.DataFrame:
        import pytesseract
        from pdf2image import convert_from_path  # type: ignore

        logger.info("  → tentando OCR…")
        images = convert_from_path(self.caminho, dpi=300)
        texto = ""
        for img in images:
            texto += pytesseract.image_to_string(img, lang="por") + "\n"
        return self._parsear_texto(texto)


# ──────────────────────────────────────────────────────────────────────────────
# Factory
# ──────────────────────────────────────────────────────────────────────────────

def criar_parser(caminho: str) -> FaturaParser:
    ext = Path(caminho).suffix.lower()
    if ext == ".pdf":
        return PDFParser(caminho)
    elif ext in (".xlsx", ".xls"):
        return ExcelParser(caminho)
    elif ext == ".csv":
        return CSVParser(caminho)
    else:
        raise ValueError(f"Formato não suportado: {ext!r}. Use PDF, XLSX ou CSV.")


# ──────────────────────────────────────────────────────────────────────────────
# Gerador de Excel
# ──────────────────────────────────────────────────────────────────────────────

def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _estilo_cabecalho(ws, n_colunas: int, linha: int = 1) -> None:
    for col in range(1, n_colunas + 1):
        cell = ws.cell(row=linha, column=col)
        cell.fill = _fill(CORES["cabecalho_bg"])
        cell.font = Font(bold=True, color=CORES["cabecalho_fg"], size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="000000"))
    ws.row_dimensions[linha].height = 22


def _estilo_dado(cell, is_par: bool, is_valor: bool) -> None:
    cell.fill = _fill(CORES["linha_par"] if is_par else CORES["linha_impar"])
    cell.alignment = Alignment(vertical="center")
    if is_valor:
        cell.number_format = FORMATO_MOEDA
        if isinstance(cell.value, (int, float)) and cell.value < 0:
            cell.font = Font(color=CORES["vermelho"])


_LARGURAS = {
    "Data": 12, "Descrição": 46, "Valor (R$)": 16,
    "Tipo": 10, "Parcela": 10, "Cartão": 20,
    "Categoria": 20, "Observações": 28,
}


def _ajustar_colunas(ws) -> None:
    for cells in ws.columns:
        header = str(cells[0].value or "")
        ws.column_dimensions[get_column_letter(cells[0].column)].width = (
            _LARGURAS.get(header, min(max(len(str(c.value or "")) for c in cells) + 4, 50))
        )


def _sheet_transacoes(wb: Workbook, df: pd.DataFrame, nome: str) -> None:
    nome = nome[:31]
    if nome in wb.sheetnames:
        del wb[nome]
    ws = wb.create_sheet(nome)

    if df.empty:
        ws.append(["Nenhuma transação encontrada."])
        return

    df_out = df.copy()
    df_out["Valor (R$)"] = pd.to_numeric(df_out["Valor (R$)"], errors="coerce")

    # Cabeçalho
    ws.append(list(df_out.columns))
    _estilo_cabecalho(ws, len(df_out.columns))

    try:
        valor_col_idx = list(df_out.columns).index("Valor (R$)") + 1
    except ValueError:
        valor_col_idx = 3

    # Dados
    for i, row in enumerate(df_out.itertuples(index=False), start=2):
        ws.append(list(row))
        is_par = i % 2 == 0
        for col in range(1, len(df_out.columns) + 1):
            _estilo_dado(ws.cell(row=i, column=col), is_par, col == valor_col_idx)
        ws.row_dimensions[i].height = 16

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    _ajustar_colunas(ws)


def _sheet_resumo(wb: Workbook, df_total: pd.DataFrame,
                  por_cartao: Dict[str, pd.DataFrame]) -> None:
    nome = "Resumo"
    if nome in wb.sheetnames:
        del wb[nome]
    ws = wb.create_sheet(nome)

    linha = 1

    # Título
    ws.merge_cells(f"A{linha}:F{linha}")
    c = ws.cell(row=linha, column=1, value="RESUMO DE FATURAS")
    c.fill = _fill(CORES["titulo_resumo"])
    c.font = Font(bold=True, color="FFFFFF", size=14)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[linha].height = 28
    linha += 1

    ws.cell(row=linha, column=1,
            value=f'Gerado em: {datetime.now().strftime("%d/%m/%Y %H:%M")}')
    ws.cell(row=linha, column=1).font = Font(italic=True, color="888888", size=10)
    linha += 2

    # Cabeçalhos da tabela
    hdrs = ["Cartão", "Qtd. Transações", "Total Débitos (R$)",
            "Total Créditos (R$)", "Total Líquido (R$)", "Período"]
    ws.append(hdrs)
    _estilo_cabecalho(ws, len(hdrs), linha)
    linha += 1

    total_deb = total_cred = total_qtd = 0.0

    for i, (cartao, df_c) in enumerate(por_cartao.items()):
        vals = pd.to_numeric(df_c["Valor (R$)"], errors="coerce").dropna()
        deb = float(vals[vals > 0].sum())
        cred = float(vals[vals < 0].sum())
        qtd = len(df_c)
        liquido = deb + cred

        try:
            datas = pd.to_datetime(df_c["Data"], format="%d/%m/%Y", errors="coerce").dropna()
            periodo = (f"{datas.min().strftime('%d/%m/%Y')} a "
                       f"{datas.max().strftime('%d/%m/%Y')}") if len(datas) else ""
        except Exception:
            periodo = ""

        row_data = [cartao, qtd, deb, cred, liquido, periodo]
        for j, v in enumerate(row_data, 1):
            cell = ws.cell(row=linha, column=j, value=v)
            cell.fill = _fill(CORES["linha_par"] if i % 2 == 0 else CORES["linha_impar"])
            cell.alignment = Alignment(vertical="center")
            if j in (3, 4, 5):
                cell.number_format = FORMATO_MOEDA
                if isinstance(v, float) and v < 0:
                    cell.font = Font(color=CORES["vermelho"])

        total_deb += deb
        total_cred += cred
        total_qtd += qtd
        linha += 1

    # Linha de totais
    totais = ["TOTAL GERAL", int(total_qtd), total_deb, total_cred,
               total_deb + total_cred, ""]
    for j, v in enumerate(totais, 1):
        cell = ws.cell(row=linha, column=j, value=v)
        cell.fill = _fill(CORES["total_bg"])
        cell.font = Font(bold=True, size=11)
        cell.border = Border(top=Side(style="medium"))
        if j in (3, 4, 5):
            cell.number_format = FORMATO_MOEDA
            if isinstance(v, float) and v < 0:
                cell.font = Font(bold=True, color=CORES["vermelho"])
    linha += 2

    # Totais por categoria (se preenchido)
    cats_df = df_total[df_total["Categoria"].notna() & (df_total["Categoria"] != "")]
    if not cats_df.empty:
        ws.cell(row=linha, column=1, value="TOTAIS POR CATEGORIA").fill = \
            _fill(CORES["titulo_resumo"])
        ws.cell(row=linha, column=1).font = Font(bold=True, color="FFFFFF")
        ws.merge_cells(f"A{linha}:C{linha}")
        linha += 1

        cats = (
            cats_df.groupby("Categoria")["Valor (R$)"]
            .apply(lambda x: pd.to_numeric(x, errors="coerce").sum())
            .reset_index()
            .sort_values("Valor (R$)", ascending=False)
        )
        for _, r in cats.iterrows():
            ws.cell(row=linha, column=1, value=r["Categoria"])
            c2 = ws.cell(row=linha, column=2, value=r["Valor (R$)"])
            c2.number_format = FORMATO_MOEDA
            if r["Valor (R$)"] < 0:
                c2.font = Font(color=CORES["vermelho"])
            linha += 1

    # Larguras fixas
    for col_letter, width in zip("ABCDEF", [22, 18, 20, 20, 20, 30]):
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "A5"


def gerar_excel(frames: List[pd.DataFrame], caminho_saida: str) -> None:
    """Cria o arquivo Excel final com todas as abas."""
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    df_total = pd.concat(frames, ignore_index=True)

    # Ordenar por data
    try:
        df_total.insert(0, "_s", pd.to_datetime(df_total["Data"], format="%d/%m/%Y", errors="coerce"))
        df_total = df_total.sort_values("_s").drop(columns=["_s"]).reset_index(drop=True)
    except Exception:
        pass

    # Agrupar por cartão
    por_cartao: Dict[str, pd.DataFrame] = {}
    for cartao in df_total["Cartão"].unique():
        por_cartao[cartao] = df_total[df_total["Cartão"] == cartao].reset_index(drop=True)

    # Abas
    _sheet_resumo(wb, df_total, por_cartao)
    _sheet_transacoes(wb, df_total, "Todas as Transações")
    for cartao, df_c in por_cartao.items():
        _sheet_transacoes(wb, df_c, cartao)

    wb.active = wb["Resumo"]
    wb.save(caminho_saida)

    # Relatório final
    vals = pd.to_numeric(df_total["Valor (R$)"], errors="coerce")
    logger.info(f"\n{'─'*50}")
    logger.info(f"✓ Arquivo salvo: {caminho_saida}")
    logger.info(f"  Total de transações : {len(df_total)}")
    logger.info(f"  Total débitos       : R$ {vals[vals > 0].sum():,.2f}")
    logger.info(f"  Total créditos      : R$ {abs(vals[vals < 0].sum()):,.2f}")
    logger.info(f"  Saldo líquido       : R$ {vals.sum():,.2f}")
    logger.info(f"{'─'*50}")
    for cartao, df_c in por_cartao.items():
        v = pd.to_numeric(df_c["Valor (R$)"], errors="coerce")
        logger.info(
            f"  {cartao:<20} {len(df_c):>3} transações | "
            f"Débitos R$ {v[v>0].sum():>10,.2f} | "
            f"Créditos R$ {abs(v[v<0].sum()):>10,.2f}"
        )


# ──────────────────────────────────────────────────────────────────────────────
# CLI
# ──────────────────────────────────────────────────────────────────────────────

def main() -> None:
    ap = argparse.ArgumentParser(
        prog="processar_faturas.py",
        description="Consolida faturas de cartão de crédito em uma planilha Excel.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Exemplos:
  python processar_faturas.py fatura_itau.pdf fatura_xp.xlsx
  python processar_faturas.py *.csv --saida fevereiro_2025.xlsx
  python processar_faturas.py fatura.pdf --cartao "Itaú Visa Gold" --verbose
        """,
    )
    ap.add_argument("arquivos", nargs="+", help="Arquivos de fatura (PDF, XLSX ou CSV)")
    ap.add_argument(
        "-s", "--saida",
        default=None,
        help="Arquivo de saída (padrão: faturas_consolidadas_AAAA-MM-DD.xlsx)",
    )
    ap.add_argument(
        "-v", "--verbose",
        action="store_true",
        help="Logs detalhados",
    )
    ap.add_argument(
        "--cartao",
        nargs="+",
        metavar="NOME",
        help="Sobrescreve o nome detectado para cada arquivo, na mesma ordem",
    )
    ap.add_argument(
        "--senha",
        metavar="SENHA",
        default=None,
        help="Senha para PDFs protegidos (ex: CPF ou data de nascimento)",
    )
    args = ap.parse_args()

    if args.verbose:
        logger.setLevel(logging.DEBUG)

    saida = args.saida or f"faturas_consolidadas_{datetime.now():%Y-%m-%d}.xlsx"

    frames: List[pd.DataFrame] = []

    for i, arq in enumerate(args.arquivos):
        p = Path(arq)
        if not p.exists():
            logger.error(f"✗ Arquivo não encontrado: {arq}")
            continue

        logger.info(f"\n[{i+1}/{len(args.arquivos)}] {p.name}")

        try:
            parser = criar_parser(str(p))

            if args.cartao and i < len(args.cartao):
                parser.nome_cartao = args.cartao[i]

            if args.senha and isinstance(parser, PDFParser):
                parser.senha = args.senha

            df = parser.parsear()

            if df.empty:
                logger.warning(f"  ⚠ Nenhuma transação extraída de {p.name}")
                continue

            logger.info(f"  ✓ {len(df)} transações extraídas")
            frames.append(df)

        except ValueError as e:
            logger.error(f"  ✗ {e}")
        except Exception as e:
            logger.error(f"  ✗ Erro inesperado em {p.name}: {e}")
            if args.verbose:
                import traceback
                traceback.print_exc()

    if not frames:
        logger.error("\n✗ Nenhuma fatura processada. Verifique os arquivos e tente novamente.")
        sys.exit(1)

    logger.info("\nGerando planilha Excel…")
    gerar_excel(frames, saida)


if __name__ == "__main__":
    main()
